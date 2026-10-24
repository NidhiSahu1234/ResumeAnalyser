"""Validated text extraction for uploaded resumes (PDF and DOCX).

Error messages are written for end users and never include file contents, so resume text
cannot leak into logs or error output through this module.
"""
from __future__ import annotations

import os
import re
from io import BytesIO

MAX_BYTES = 5 * 1024 * 1024
MAX_PDF_PAGES = 15
MIN_TEXT_CHARS = 40


class ResumeFileError(ValueError):
    """Raised for any problem with an uploaded resume. ``str(error)`` is safe to show users."""


def normalize_text(text: str) -> str:
    """Tidy extracted text: unify bullets and line breaks, drop control characters."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ").replace("", "•")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ResumeFileError("This PDF is password-protected. Remove the password or paste the text instead.")
        pages = reader.pages[:MAX_PDF_PAGES]
        return "\n".join((page.extract_text() or "") for page in pages)
    except ResumeFileError:
        raise
    except Exception as exc:  # pypdf raises many types for damaged files
        raise ResumeFileError(
            "This PDF could not be read. It may be damaged. Try re-exporting it or paste the text instead.") from exc


def _docx_text(data: bytes) -> str:
    try:
        from docx import Document

        document = Document(BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:  # many resume templates put content in tables
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)
    except Exception as exc:
        raise ResumeFileError(
            "This Word file could not be read. It may be damaged or not a real .docx. "
            "Try re-saving it or paste the text instead.") from exc


def extract_resume_text(filename: str, data: bytes) -> str:
    """Return clean text from an uploaded resume or raise :class:`ResumeFileError`."""
    if not data:
        raise ResumeFileError("The uploaded file is empty.")
    if len(data) > MAX_BYTES:
        raise ResumeFileError(f"The file is larger than {MAX_BYTES // (1024 * 1024)} MB. Upload a smaller file.")

    extension = os.path.splitext(filename or "")[1].lower()
    if extension == ".pdf":
        if not data.lstrip()[:5] == b"%PDF-":
            raise ResumeFileError("This file has a .pdf extension but is not a valid PDF.")
        text = _pdf_text(data)
    elif extension == ".docx":
        if not data[:4] == b"PK\x03\x04":
            raise ResumeFileError("This file has a .docx extension but is not a valid Word document.")
        text = _docx_text(data)
    else:
        raise ResumeFileError("Unsupported file type. Upload a PDF or DOCX file, or paste the text instead.")

    text = normalize_text(text)
    if len(text) < MIN_TEXT_CHARS:
        raise ResumeFileError(
            "No readable text was found. Scanned or image-only files need OCR; "
            "paste the resume text instead.")
    return text
