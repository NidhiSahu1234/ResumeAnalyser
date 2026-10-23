from io import BytesIO

import pytest
from docx import Document

from matching import resume_io as rio

SECRET = "SECRET-RESUME-MARKER-12345"


def make_pdf(text: str) -> bytes:
    """Smallest valid single-page PDF with one line of text (no extra dependencies)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def make_docx(paragraphs, table_cells=()):
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_cells:
        table = document.add_table(rows=1, cols=len(table_cells))
        for cell, value in zip(table.rows[0].cells, table_cells):
            cell.text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


LONG_LINE = "Backend developer with Python Flask and Docker experience in payments"


def test_reads_valid_pdf():
    text = rio.extract_resume_text("resume.pdf", make_pdf(LONG_LINE))
    assert "Python Flask and Docker" in text


def test_reads_valid_docx_including_tables():
    data = make_docx([LONG_LINE, "Second paragraph about projects and education"], table_cells=("Skills", "Python, SQL"))
    text = rio.extract_resume_text("RESUME.DOCX", data)  # extension check is case-insensitive
    assert "Backend developer" in text and "Python, SQL" in text


@pytest.mark.parametrize("name,data,fragment", [
    ("resume.pdf", b"", "empty"),
    ("resume.txt", b"plain text resume " * 10, "Unsupported"),
    ("resume", b"x" * 100, "Unsupported"),
    (None, b"x" * 100, "Unsupported"),
    ("resume.pdf", b"this is not a pdf at all, just text " * 5, "not a valid PDF"),
    ("resume.docx", b"this is not a zip, just text " * 5, "not a valid Word"),
    ("resume.exe", b"MZ" + b"\x00" * 100, "Unsupported"),
])
def test_rejects_bad_files_with_clear_messages(name, data, fragment):
    with pytest.raises(rio.ResumeFileError, match=fragment):
        rio.extract_resume_text(name, data)


def test_rejects_oversized_files():
    with pytest.raises(rio.ResumeFileError, match="larger than"):
        rio.extract_resume_text("resume.pdf", b"%PDF-1.4" + b"0" * rio.MAX_BYTES)


def test_corrupt_pdf_and_docx_fail_gracefully_without_leaking_content():
    corrupt_pdf = b"%PDF-1.4\n" + SECRET.encode() + b"\n\x00garbage" * 20
    corrupt_docx = b"PK\x03\x04" + SECRET.encode() + b"\x00garbage" * 20
    for name, data in (("a.pdf", corrupt_pdf), ("b.docx", corrupt_docx)):
        with pytest.raises(rio.ResumeFileError) as caught:
            rio.extract_resume_text(name, data)
        assert SECRET not in str(caught.value)


def test_image_only_or_blank_documents_are_reported_not_returned_empty():
    with pytest.raises(rio.ResumeFileError, match="No readable text"):
        rio.extract_resume_text("blank.pdf", make_pdf(" "))
    with pytest.raises(rio.ResumeFileError, match="No readable text"):
        rio.extract_resume_text("blank.docx", make_docx(["", "  "]))


def test_normalize_text_cleans_bullets_and_control_characters():
    cleaned = rio.normalize_text(" Built it\x00\r\n\r\n\r\n\r\nNext\xa0line   \n")
    assert cleaned == "• Built it\n\nNext line"
